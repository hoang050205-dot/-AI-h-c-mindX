import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import sys

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Target file path (strictly inside sample-data)
WORKSPACE_DIR = r"c:\Minh Hoang\Antigravity học\my-workspace"
FILE_SAMPLE_DATA = os.path.join(WORKSPACE_DIR, "sample-data", "marketing_campaigns.xlsx")

# Create workbook
wb = openpyxl.Workbook()

# Sheet 1: Main Data (70 rows, 10 columns)
ws = wb.active
ws.title = "Marketing_Campaigns"

headers = [
    "Campaign_ID",
    "Campaign_Name",
    "Channel",
    "Budget",
    "Currency",
    "Status",
    "Start_Date",
    "End_Date",
    "Actual_Spend",
    "Manager"
]

data = [
    # --- BATCH 1: 20 dòng ban đầu (CMP001 - CMP020) ---
    ["CMP001", "Summer Sale Mega Live", "Facebook Ads", 45000000, "VND", "Completed", "2026-06-01", "2026-06-30", 44500000, "Nguyễn Văn An"],
    ["CMP002", "Google Search Brand Protection", "Google Ads", 2500, "USD", "Active", "2026-07-01", "2026-08-31", 1850, "Trần Thị Mai"],
    ["CMP003", "TikTok Dance Challenge GenZ", "TikTok Ads", 80000000, "VND", "Active", "2026-08-15", "Sau lễ", 42000000, "Lê Hoàng Nam"],
    ["CMP004", "SEO Content Pillar Q3", "Organic Search", 15000000, None, "Active", "2026-07-01", "2026-09-30", 0, "Phạm Thu Trang"],
    ["CMP005", "LinkedIn B2B Lead Summit", "LinkedIn Ads", 3200, "USD", "Completed", "2026-05-01", "2026-05-31", 3180, "Đặng Minh Quân"],
    ["CMP006", "National TVC & Highway Billboard", "OOH / TVC", 1850000000, "VND", "Active", "2026-09-01", "2026-11-30", 620000000, "Nguyễn Văn An"],
    ["CMP007", "Facebook Retargeting Catalog", "Facebook Ads", 30000000, None, "Active", "2026-08-01", "2026-08-31", 28500000, "Trần Thị Mai"],
    ["CMP008", "YouTube Tech Review Influencer", "YouTube", 4000, "USD", "Active", "Tháng sau", "2026-10-31", 0, "Lê Hoàng Nam"],
    ["CMP009", "Email Drip Welcome New Users", "Email Marketing", 12000000, "VND", "Completed", "2026-04-01", "2026-06-30", 11800000, "Phạm Thu Trang"],
    ["CMP010", "Shopee 9.9 Super Shopping Day", "Shopee Ads", 65000000, None, "Completed", "2026-09-01", "2026-09-10", 64200000, "Đặng Minh Quân"],
    ["CMP011", "Mega Year-End Omnichannel Push", "Multi-channel", 3200000000, "VND", "Active", "2026-10-01", "2026-12-31", 450000000, "Nguyễn Văn An"],
    ["CMP012", "Google Performance Max Ecom", "Google Ads", 1800, "USD", "Active", "2026-08-01", "2026-08-31", 1420, "Trần Thị Mai"],
    ["CMP013", "TikTok Livestream Flash Voucher", "TikTok Ads", 50000000, "VND", "Completed", "2026-07-01", "Q3", 49000000, "Lê Hoàng Nam"],
    ["CMP014", "Mobile App Install Universal", "Google Ads", 5500, "USD", "Active", "2026-09-01", "2026-10-15", 2100, "Phạm Thu Trang"],
    ["CMP015", "Local Mall Experiential Booth", "Activation / Event", 95000000, None, "Active", "Đầu tuần sau", "2026-09-25", 0, "Đặng Minh Quân"],
    ["CMP016", "Affiliate Referral Commission", "Affiliate Network", 20000000, None, "Paused", "2026-06-15", "2026-07-15", 8500000, "Nguyễn Văn An"],
    ["CMP017", "Community Fanpage Engagement Minigame", "Social Media", 18000000, "VND", "Active", "2026-08-01", "2026-08-31", 17200000, "Trần Thị Mai"],
    ["CMP018", "Podcast Sponsorship Top Tech VN", "Audio Network", 35000000, "VND", "Planned", "2026-11-01", "2026-11-30", 0, "Lê Hoàng Nam"],
    ["CMP019", "SMS Brandname Flash Promo", "SMS Marketing", 22000000, "VND", "Completed", "2026-05-15", "2026-05-17", 21900000, "Phạm Thu Trang"],
    ["CMP020", "Lead Magnet Industry Report 2026", "LinkedIn Ads", 28000000, "VND", "Active", "2026-07-15", "2026-08-15", 27600000, "Đặng Minh Quân"],

    # --- BATCH 2: 50 dòng bổ sung (CMP021 - CMP070) ---
    ["CMP021", "TikTok Spark Ads Conversion Boost", "TikTok Ads", 60000000, "VND", "Completed", "2026-06-10", "2026-07-10", 59400000, "Lê Hoàng Nam"],
    # USD 1/5 new
    ["CMP022", "Google Display Network Retargeting", "Google Ads", 3500, "USD", "Active", "2026-08-01", "2026-09-15", 2750, "Trần Thị Mai"],
    # Date 1/5 new: "Sau Tết"
    ["CMP023", "Spring Festive Brand Giveaway", "Facebook Ads", 35000000, "VND", "Completed", "2026-01-20", "Sau Tết", 34800000, "Nguyễn Văn An"],
    ["CMP024", "Shopee Flash Voucher Hunt 10.10", "Shopee Ads", 45000000, "VND", "Planned", "2026-10-05", "2026-10-10", 0, "Hoàng Đức Thắng"],
    # No unit 1/5 new
    ["CMP025", "SEO Link Building Authority Surge", "Organic Search", 40000000, None, "Active", "2026-07-01", "2026-09-30", 38200000, "Phạm Thu Trang"],
    # Budget > 1 tỷ 1/5 new
    ["CMP026", "TVC Tet Countdown Prime Time", "OOH / TVC", 1200000000, "VND", "Active", "2026-09-15", "2026-12-31", 350000000, "Đặng Minh Quân"],
    # Active Spend=0 1/5 new
    ["CMP027", "LinkedIn Enterprise Account ABM", "LinkedIn Ads", 50000000, "VND", "Active", "2026-08-01", "2026-08-31", 0, "Vũ Thị Bích Ngọc"],
    ["CMP028", "Micro-Influencer Foodie Review", "TikTok Ads", 25000000, "VND", "Completed", "2026-05-01", "2026-05-20", 24600000, "Lê Hoàng Nam"],
    ["CMP029", "Zalo ZNS Order Status Automation", "Zalo Ads", 16000000, "VND", "Active", "2026-08-10", "2026-09-10", 15300000, "Phạm Thu Trang"],
    # USD 2/5 new
    ["CMP030", "YouTube Bumper Ads Quick Reach", "YouTube", 1200, "USD", "Active", "2026-08-01", "2026-08-20", 950, "Trần Thị Mai"],
    ["CMP031", "Lazada 11.11 Mega Pre-Hype", "Lazada Ads", 55000000, "VND", "Planned", "2026-11-01", "2026-11-11", 0, "Hoàng Đức Thắng"],
    # Date 2/5 new: "Giữa tháng này"
    ["CMP032", "Facebook Lead Form Test Drive", "Facebook Ads", 28000000, "VND", "Active", "Giữa tháng này", "2026-09-30", 12400000, "Nguyễn Văn An"],
    ["CMP033", "Email Re-engagement Dormant Leads", "Email Marketing", 14000000, "VND", "Completed", "2026-06-01", "2026-06-30", 13700000, "Phạm Thu Trang"],
    # Active Spend=0 2/5 new
    ["CMP034", "University Tour Activation Booth", "Activation / Event", 42000000, "VND", "Active", "2026-09-01", "2026-09-30", 0, "Đặng Minh Quân"],
    # No unit 2/5 new
    ["CMP035", "TikTok Brand Takeover TopView", "TikTok Ads", 85000000, None, "Completed", "2026-07-20", "2026-07-22", 84900000, "Lê Hoàng Nam"],
    # Budget > 1 tỷ 2/5 new
    ["CMP036", "National Roadshow 10 Provinces", "Activation / Event", 2500000000, "VND", "Active", "2026-08-01", "2026-11-15", 890000000, "Đặng Minh Quân"],
    ["CMP037", "Google App Install iOS Specific", "Google Ads", 38000000, "VND", "Active", "2026-08-15", "2026-09-15", 36400000, "Trần Thị Mai"],
    ["CMP038", "Shopee Live Affiliate KOC Pool", "Shopee Ads", 48000000, "VND", "Completed", "2026-06-15", "2026-07-15", 47800000, "Hoàng Đức Thắng"],
    ["CMP039", "PR Guest Post Tech News Portals", "Influencer / KOL", 30000000, "VND", "Completed", "2026-04-10", "2026-05-10", 30000000, "Vũ Thị Bích Ngọc"],
    ["CMP040", "Facebook Carousel Video Ads", "Facebook Ads", 22000000, "VND", "Paused", "2026-05-20", "2026-06-20", 11500000, "Nguyễn Văn An"],
    # USD 3/5 new
    ["CMP041", "Global Tech Expo Booth Digital Push", "Google Ads", 6000, "USD", "Active", "2026-09-01", "2026-10-15", 4100, "Trần Thị Mai"],
    # Date 3/5 new: "Cuối Q3"
    ["CMP042", "Mid-Autumn Festival Gift Box Promo", "Social Media", 32000000, "VND", "Active", "2026-08-01", "Cuối Q3", 26000000, "Lê Hoàng Nam"],
    ["CMP043", "Loyalty Tier Upgrade SMS Push", "SMS Marketing", 18000000, "VND", "Completed", "2026-07-01", "2026-07-03", 17900000, "Phạm Thu Trang"],
    # Active Spend=0 3/5 new
    ["CMP044", "Podcast Mid-roll Host Read Sponsorship", "Audio Network", 28000000, "VND", "Active", "2026-07-15", "2026-08-15", 0, "Vũ Thị Bích Ngọc"],
    # No unit 3/5 new
    ["CMP045", "Community Discord Server Onboarding", "Social Media", 18000000, None, "Active", "2026-08-01", "2026-08-31", 16500000, "Lê Hoàng Nam"],
    # Budget > 1 tỷ 3/5 new
    ["CMP046", "Brand Rebranding Mega Launch Event", "Multi-channel", 1450000000, "VND", "Active", "2026-09-01", "2026-11-30", 480000000, "Đặng Minh Quân"],
    ["CMP047", "Google Shopping Feed Optimization", "Google Ads", 26000000, "VND", "Completed", "2026-05-15", "2026-06-15", 25700000, "Trần Thị Mai"],
    ["CMP048", "Shopee In-app Banner Carousel", "Shopee Ads", 34000000, "VND", "Active", "2026-08-20", "2026-09-20", 31800000, "Hoàng Đức Thắng"],
    ["CMP049", "TikTok Live Shopping Creator Battle", "TikTok Ads", 52000000, "VND", "Completed", "2026-06-01", "2026-06-15", 51500000, "Lê Hoàng Nam"],
    ["CMP050", "Affiliate Coupon Site Integration", "Affiliate Network", 21000000, "VND", "Active", "2026-08-01", "2026-08-31", 19800000, "Nguyễn Văn An"],
    ["CMP051", "Zalo Mini App User Acquisition", "Zalo Ads", 30000000, "VND", "Completed", "2026-06-15", "2026-07-15", 29300000, "Phạm Thu Trang"],
    # USD 4/5 new
    ["CMP052", "YouTube TrueView For Action", "YouTube", 2800, "USD", "Active", "2026-08-10", "2026-09-20", 2100, "Trần Thị Mai"],
    # Date 4/5 new: "Tuần tới"
    ["CMP053", "Flash Clearance Sale Weekend", "Facebook Ads", 24000000, "VND", "Active", "Tuần tới", "2026-09-30", 8900000, "Nguyễn Văn An"],
    ["CMP054", "E-Newsletter Monthly Tech Insights", "Email Marketing", 11000000, "VND", "Completed", "2026-05-01", "2026-05-31", 10800000, "Phạm Thu Trang"],
    # No unit 4/5 new
    ["CMP055", "TikTok Sound Branding Audio Logo", "TikTok Ads", 55000000, None, "Active", "2026-07-01", "2026-08-15", 53200000, "Lê Hoàng Nam"],
    # Active Spend=0 4/5 new
    ["CMP056", "Airport VIP Lounge Digital Screen", "OOH / TVC", 65000000, "VND", "Active", "2026-08-20", "2026-09-20", 0, "Đặng Minh Quân"],
    # Budget > 1 tỷ 4/5 new
    ["CMP057", "Global Music Festival Co-sponsorship", "Multi-channel", 4000000000, "VND", "Active", "2026-10-01", "2026-12-31", 950000000, "Đặng Minh Quân"],
    ["CMP058", "Shopee Brand Exclusive Livestream", "Shopee Ads", 42000000, "VND", "Completed", "2026-06-25", "2026-06-28", 41600000, "Hoàng Đức Thắng"],
    ["CMP059", "Lead Magnet B2B SaaS Benchmark", "LinkedIn Ads", 31000000, "VND", "Active", "2026-08-01", "2026-08-31", 29400000, "Vũ Thị Bích Ngọc"],
    ["CMP060", "Tech Blog Guest Article Backlinks", "Organic Search", 19000000, "VND", "Completed", "2026-04-15", "2026-05-15", 18700000, "Phạm Thu Trang"],
    ["CMP061", "Facebook Lookalike High LTV Push", "Facebook Ads", 36000000, "VND", "Active", "2026-08-15", "2026-09-15", 34900000, "Nguyễn Văn An"],
    # Date 5/5 new: "Hết mùa hè"
    ["CMP062", "Summer Travel Gear Promo", "TikTok Ads", 46000000, "VND", "Completed", "2026-06-01", "Hết mùa hè", 45500000, "Lê Hoàng Nam"],
    # USD 5/5 new
    ["CMP063", "Google Discovery Feed Carousel", "Google Ads", 4200, "USD", "Active", "2026-08-15", "2026-09-30", 3450, "Trần Thị Mai"],
    ["CMP064", "Lazada Brand Mega Offer 12.12", "Lazada Ads", 58000000, "VND", "Planned", "2026-12-01", "2026-12-12", 0, "Hoàng Đức Thắng"],
    # No unit 5/5 new
    ["CMP065", "Retail Store Window Decal Branding", "Activation / Event", 32000000, None, "Completed", "2026-05-10", "2026-05-25", 31500000, "Đặng Minh Quân"],
    ["CMP066", "Automated Birthday Coupon Drip", "Email Marketing", 15000000, "VND", "Active", "2026-07-01", "2026-12-31", 6800000, "Phạm Thu Trang"],
    # Active Spend=0 5/5 new
    ["CMP067", "Gaming Streamer Product Placement", "YouTube", 55000000, "VND", "Active", "2026-09-10", "2026-10-10", 0, "Vũ Thị Bích Ngọc"],
    # Budget > 1 tỷ 5/5 new
    ["CMP068", "Nationwide LED City Screen Takeover", "OOH / TVC", 1750000000, "VND", "Active", "2026-09-01", "2026-11-30", 520000000, "Đặng Minh Quân"],
    ["CMP069", "TikTok Duet Challenge With Celeb", "TikTok Ads", 62000000, "VND", "Active", "2026-08-25", "2026-09-25", 58700000, "Lê Hoàng Nam"],
    ["CMP070", "Year-End Thank You Customer Gala", "Activation / Event", 88000000, "VND", "Planned", "2026-12-20", "2026-12-25", 0, "Đặng Minh Quân"],
]

# Write header
ws.append(headers)

# Styling definitions
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color="E2E8F0"),
    right=Side(style='thin', color="E2E8F0"),
    top=Side(style='thin', color="E2E8F0"),
    bottom=Side(style='thin', color="E2E8F0")
)

data_font = Font(name="Segoe UI", size=10, color="0F172A")
zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

# Apply header styles
ws.row_dimensions[1].height = 26
for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# Write and style data rows
for row_idx, row_data in enumerate(data, start=2):
    ws.append(row_data)
    ws.row_dimensions[row_idx].height = 20
    is_even = (row_idx % 2 == 0)
    current_fill = zebra_fill if is_even else white_fill

    for col_idx in range(1, len(row_data) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font = data_font
        cell.fill = current_fill
        cell.border = thin_border

        # Alignment & Number format based on column
        col_name = headers[col_idx - 1]
        if col_name in ["Campaign_ID", "Currency", "Status", "Start_Date", "End_Date"]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_name in ["Budget", "Actual_Spend"]:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

# Specific column width fine-tuning
ws.column_dimensions["A"].width = 15  # Campaign_ID
ws.column_dimensions["B"].width = 40  # Campaign_Name
ws.column_dimensions["C"].width = 22  # Channel
ws.column_dimensions["D"].width = 18  # Budget
ws.column_dimensions["E"].width = 12  # Currency
ws.column_dimensions["F"].width = 14  # Status
ws.column_dimensions["G"].width = 16  # Start_Date
ws.column_dimensions["H"].width = 16  # End_Date
ws.column_dimensions["I"].width = 18  # Actual_Spend
ws.column_dimensions["J"].width = 22  # Manager

# Freeze top row
ws.freeze_panes = "A2"
ws.views.sheetView[0].showGridLines = True

# Sheet 2: Audit Guide (Tất cả các case lỗi chi tiết)
ws_audit = wb.create_sheet(title="Audit_Guide")
ws_audit.views.sheetView[0].showGridLines = True

audit_headers = ["Category", "Batch", "Row_Index", "Campaign_ID", "Campaign_Name", "Flaw_Description", "Audit_Detail"]
ws_audit.append(audit_headers)
ws_audit.row_dimensions[1].height = 26

for col_idx in range(1, len(audit_headers) + 1):
    cell = ws_audit.cell(row=1, column=col_idx)
    cell.font = header_font
    cell.fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    cell.alignment = header_align
    cell.border = thin_border

audit_rows = [
    # --- Tiền tệ USD (Tổng 10 case: 5 cũ + 5 mới) ---
    ["Tiền tệ USD", "Batch 1", 2, "CMP002", "Google Search Brand Protection", "Budget tính bằng USD ($2,500)", "Currency='USD', Actual_Spend=1,850"],
    ["Tiền tệ USD", "Batch 1", 5, "CMP005", "LinkedIn B2B Lead Summit", "Budget tính bằng USD ($3,200)", "Currency='USD', Actual_Spend=3,180"],
    ["Tiền tệ USD", "Batch 1", 8, "CMP008", "YouTube Tech Review Influencer", "Budget tính bằng USD ($4,000)", "Currency='USD', Actual_Spend=0"],
    ["Tiền tệ USD", "Batch 1", 12, "CMP012", "Google Performance Max Ecom", "Budget tính bằng USD ($1,800)", "Currency='USD', Actual_Spend=1,420"],
    ["Tiền tệ USD", "Batch 1", 14, "CMP014", "Mobile App Install Universal", "Budget tính bằng USD ($5,500)", "Currency='USD', Actual_Spend=2,100"],
    ["Tiền tệ USD", "Batch 2", 22, "CMP022", "Google Display Network Retargeting", "Budget tính bằng USD ($3,500)", "Currency='USD', Actual_Spend=2,750"],
    ["Tiền tệ USD", "Batch 2", 30, "CMP030", "YouTube Bumper Ads Quick Reach", "Budget tính bằng USD ($1,200)", "Currency='USD', Actual_Spend=950"],
    ["Tiền tệ USD", "Batch 2", 41, "CMP041", "Global Tech Expo Booth Digital Push", "Budget tính bằng USD ($6,000)", "Currency='USD', Actual_Spend=4,100"],
    ["Tiền tệ USD", "Batch 2", 52, "CMP052", "YouTube TrueView For Action", "Budget tính bằng USD ($2,800)", "Currency='USD', Actual_Spend=2,100"],
    ["Tiền tệ USD", "Batch 2", 63, "CMP063", "Google Discovery Feed Carousel", "Budget tính bằng USD ($4,200)", "Currency='USD', Actual_Spend=3,450"],

    # --- Không đơn vị tiền tệ (Tổng 10 case: 5 cũ + 5 mới) ---
    ["Không đơn vị tiền tệ", "Batch 1", 4, "CMP004", "SEO Content Pillar Q3", "Chỉ có số 15,000,000, thiếu đơn vị", "Currency=None / Trống"],
    ["Không đơn vị tiền tệ", "Batch 1", 7, "CMP007", "Facebook Retargeting Catalog", "Chỉ có số 30,000,000, thiếu đơn vị", "Currency=None / Trống"],
    ["Không đơn vị tiền tệ", "Batch 1", 10, "CMP010", "Shopee 9.9 Super Shopping Day", "Chỉ có số 65,000,000, thiếu đơn vị", "Currency=None / Trống"],
    ["Không đơn vị tiền tệ", "Batch 1", 15, "CMP015", "Local Mall Experiential Booth", "Chỉ có số 95,000,000, thiếu đơn vị", "Currency=None / Trống"],
    ["Không đơn vị tiền tệ", "Batch 1", 16, "CMP016", "Affiliate Referral Commission", "Chỉ có số 20,000,000, thiếu đơn vị", "Currency=None / Trống"],
    ["Không đơn vị tiền tệ", "Batch 2", 25, "CMP025", "SEO Link Building Authority Surge", "Chỉ có số 40,000,000, thiếu đơn vị", "Currency=None / Trống"],
    ["Không đơn vị tiền tệ", "Batch 2", 35, "CMP035", "TikTok Brand Takeover TopView", "Chỉ có số 85,000,000, thiếu đơn vị", "Currency=None / Trống"],
    ["Không đơn vị tiền tệ", "Batch 2", 45, "CMP045", "Community Discord Server Onboarding", "Chỉ có số 18,000,000, thiếu đơn vị", "Currency=None / Trống"],
    ["Không đơn vị tiền tệ", "Batch 2", 55, "CMP055", "TikTok Sound Branding Audio Logo", "Chỉ có số 55,000,000, thiếu đơn vị", "Currency=None / Trống"],
    ["Không đơn vị tiền tệ", "Batch 2", 65, "CMP065", "Retail Store Window Decal Branding", "Chỉ có số 32,000,000, thiếu đơn vị", "Currency=None / Trống"],

    # --- Active nhưng Spend = 0 (Tổng 8 case: 3 cũ + 5 mới) ---
    ["Active nhưng Spend = 0", "Batch 1", 4, "CMP004", "SEO Content Pillar Q3", "Trạng thái Active nhưng Actual_Spend = 0", "Status='Active', Actual_Spend=0"],
    ["Active nhưng Spend = 0", "Batch 1", 8, "CMP008", "YouTube Tech Review Influencer", "Trạng thái Active nhưng Actual_Spend = 0", "Status='Active', Actual_Spend=0"],
    ["Active nhưng Spend = 0", "Batch 1", 15, "CMP015", "Local Mall Experiential Booth", "Trạng thái Active nhưng Actual_Spend = 0", "Status='Active', Actual_Spend=0"],
    ["Active nhưng Spend = 0", "Batch 2", 27, "CMP027", "LinkedIn Enterprise Account ABM", "Trạng thái Active nhưng Actual_Spend = 0", "Status='Active', Actual_Spend=0"],
    ["Active nhưng Spend = 0", "Batch 2", 34, "CMP034", "University Tour Activation Booth", "Trạng thái Active nhưng Actual_Spend = 0", "Status='Active', Actual_Spend=0"],
    ["Active nhưng Spend = 0", "Batch 2", 44, "CMP044", "Podcast Mid-roll Host Read Sponsorship", "Trạng thái Active nhưng Actual_Spend = 0", "Status='Active', Actual_Spend=0"],
    ["Active nhưng Spend = 0", "Batch 2", 56, "CMP056", "Airport VIP Lounge Digital Screen", "Trạng thái Active nhưng Actual_Spend = 0", "Status='Active', Actual_Spend=0"],
    ["Active nhưng Spend = 0", "Batch 2", 67, "CMP067", "Gaming Streamer Product Placement", "Trạng thái Active nhưng Actual_Spend = 0", "Status='Active', Actual_Spend=0"],

    # --- Ngày tháng tự nhiên (Tổng 9 case: 4 cũ + 5 mới) ---
    ["Ngày tháng tự nhiên", "Batch 1", 3, "CMP003", "TikTok Dance Challenge GenZ", "End_Date = 'Sau lễ'", "Dữ liệu phi cấu trúc"],
    ["Ngày tháng tự nhiên", "Batch 1", 8, "CMP008", "YouTube Tech Review Influencer", "Start_Date = 'Tháng sau'", "Dữ liệu phi cấu trúc"],
    ["Ngày tháng tự nhiên", "Batch 1", 13, "CMP013", "TikTok Livestream Flash Voucher", "End_Date = 'Q3'", "Dữ liệu phi cấu trúc"],
    ["Ngày tháng tự nhiên", "Batch 1", 15, "CMP015", "Local Mall Experiential Booth", "Start_Date = 'Đầu tuần sau'", "Dữ liệu phi cấu trúc"],
    ["Ngày tháng tự nhiên", "Batch 2", 23, "CMP023", "Spring Festive Brand Giveaway", "End_Date = 'Sau Tết'", "Dữ liệu phi cấu trúc"],
    ["Ngày tháng tự nhiên", "Batch 2", 32, "CMP032", "Facebook Lead Form Test Drive", "Start_Date = 'Giữa tháng này'", "Dữ liệu phi cấu trúc"],
    ["Ngày tháng tự nhiên", "Batch 2", 42, "CMP042", "Mid-Autumn Festival Gift Box Promo", "End_Date = 'Cuối Q3'", "Dữ liệu phi cấu trúc"],
    ["Ngày tháng tự nhiên", "Batch 2", 53, "CMP053", "Flash Clearance Sale Weekend", "Start_Date = 'Tuần tới'", "Dữ liệu phi cấu trúc"],
    ["Ngày tháng tự nhiên", "Batch 2", 62, "CMP062", "Summer Travel Gear Promo", "End_Date = 'Hết mùa hè'", "Dữ liệu phi cấu trúc"],

    # --- Budget bất thường > 1 tỷ VND (Tổng 7 case: 2 cũ + 5 mới) ---
    ["Budget bất thường > 1 tỷ VND", "Batch 1", 6, "CMP006", "National TVC & Highway Billboard", "Budget = 1,850,000,000 VND", "Outlier > 1 tỷ VND"],
    ["Budget bất thường > 1 tỷ VND", "Batch 1", 11, "CMP011", "Mega Year-End Omnichannel Push", "Budget = 3,200,000,000 VND", "Outlier > 1 tỷ VND"],
    ["Budget bất thường > 1 tỷ VND", "Batch 2", 26, "CMP026", "TVC Tet Countdown Prime Time", "Budget = 1,200,000,000 VND", "Outlier > 1 tỷ VND"],
    ["Budget bất thường > 1 tỷ VND", "Batch 2", 36, "CMP036", "National Roadshow 10 Provinces", "Budget = 2,500,000,000 VND", "Outlier > 1 tỷ VND"],
    ["Budget bất thường > 1 tỷ VND", "Batch 2", 46, "CMP046", "Brand Rebranding Mega Launch Event", "Budget = 1,450,000,000 VND", "Outlier > 1 tỷ VND"],
    ["Budget bất thường > 1 tỷ VND", "Batch 2", 57, "CMP057", "Global Music Festival Co-sponsorship", "Budget = 4,000,000,000 VND", "Outlier > 1 tỷ VND"],
    ["Budget bất thường > 1 tỷ VND", "Batch 2", 68, "CMP068", "Nationwide LED City Screen Takeover", "Budget = 1,750,000,000 VND", "Outlier > 1 tỷ VND"],
]

for row_idx, r_data in enumerate(audit_rows, start=2):
    ws_audit.append(r_data)
    ws_audit.row_dimensions[row_idx].height = 20
    is_even = (row_idx % 2 == 0)
    current_fill = zebra_fill if is_even else white_fill
    for col_idx in range(1, len(r_data) + 1):
        cell = ws_audit.cell(row=row_idx, column=col_idx)
        cell.font = data_font
        cell.fill = current_fill
        cell.border = thin_border
        if col_idx in [1, 2, 3, 4]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

ws_audit.column_dimensions["A"].width = 28
ws_audit.column_dimensions["B"].width = 12
ws_audit.column_dimensions["C"].width = 12
ws_audit.column_dimensions["D"].width = 16
ws_audit.column_dimensions["E"].width = 38
ws_audit.column_dimensions["F"].width = 45
ws_audit.column_dimensions["G"].width = 45
ws_audit.freeze_panes = "A2"

# Save workbook strictly to sample-data
wb.save(FILE_SAMPLE_DATA)
print(f"Successfully saved 70 rows to: {FILE_SAMPLE_DATA}")
