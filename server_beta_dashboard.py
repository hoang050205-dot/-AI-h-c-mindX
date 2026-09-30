#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Beta Solutions BI Dashboard Server - Realtime Executive Budget Analytics
Công ty Cổ phần Beta Solutions
Chức năng:
  - Phục vụ Web Dashboard Glassmorphism Dark Mode trên cổng 8088 (hoặc tùy chọn)
  - Cung cấp API /api/raw_data và /api/data trả về dữ liệu thô (raw) từ TH_ngan_sach_phong_ban.xlsx
  - Frontend tự đảm nhiệm 100% logic lọc và tổng hợp chỉ số
  - Hỗ trợ Polling 2 giây từ client, tự động phát hiện file thay đổi (mtime)
  - Tự động mở Google Chrome khi khởi động
"""

import os
import sys

# Đảm bảo in tiếng Việt trên console Windows không bị lỗi encoding
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

import json
import time
import socket
import threading
import webbrowser
import subprocess
from datetime import datetime
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse

try:
    import openpyxl
except ImportError:
    print("LỖI: Chưa cài đặt openpyxl. Vui lòng chạy: pip install openpyxl")
    sys.exit(1)

# Đường dẫn cơ sở của Workspace
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REPORT_HTML_PATH = os.path.join(BASE_DIR, "outputs", "reports", "beta_solutions_bi_dashboard.html")

# Danh sách ứng viên file dữ liệu ngân sách
DATA_FILE_CANDIDATES = [
    os.path.join(BASE_DIR, "sample-data", "TH_ngan_sach_phong_ban.xlsx"),
    os.path.join(BASE_DIR, "TH_ngan_sach_phong_ban.xlsx")
]

# Cache dữ liệu trong RAM để polling 2s không bị nghẽn I/O
DATA_CACHE = {
    "file_path": None,
    "last_mtime": 0,
    "records": None,
    "cached_payload": None
}

def find_budget_data_file():
    """Tìm file dữ liệu TH_ngan_sach_phong_ban.xlsx theo độ ưu tiên."""
    for path in DATA_FILE_CANDIDATES:
        if os.path.exists(path):
            return path
    return None

def load_raw_records_from_excel(file_path):
    """Đọc dữ liệu thô (raw records) từ file Excel."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    headers = [ws.cell(row=1, column=j).value for j in range(1, ws.max_column + 1)]
    records = []
    
    for i in range(2, ws.max_row + 1):
        row_dict = {}
        for j, h in enumerate(headers, 1):
            val = ws.cell(row=i, column=j).value
            row_dict[h] = val
        if row_dict.get("Ma_Giao_Dich"):
            records.append(row_dict)
            
    wb.close()
    return records

def get_latest_raw_data():
    """Lấy dữ liệu thô có cơ chế so khớp mtime để tối ưu hiệu năng phản hồi < 2ms."""
    file_path = find_budget_data_file()
    if not file_path:
        return {
            "status": "error",
            "message": "Không tìm thấy file TH_ngan_sach_phong_ban.xlsx trong sample-data/.",
            "timestamp": datetime.now().strftime("%H:%M:%S")
        }

    current_mtime = os.path.getmtime(file_path)
    if (DATA_CACHE["records"] is None or 
        DATA_CACHE["file_path"] != file_path or 
        DATA_CACHE["last_mtime"] != current_mtime):
        
        print(f"[CACHE UPDATE] Phát hiện file Excel thay đổi ({os.path.basename(file_path)} - mtime: {current_mtime}). Tiến hành re-parse...")
        records = load_raw_records_from_excel(file_path)
        
        # Đếm thống kê sơ bộ
        total_ns = sum(r.get("Ngan_Sach_Duyet", 0) or 0 for r in records)
        total_ct = sum(r.get("Chi_Tieu_Thuc_Te", 0) or 0 for r in records)
        vuot_count = sum(1 for r in records if r.get("Trang_Thai") == "Vuot Ngan Sach")
        
        payload = {
            "status": "success",
            "timestamp": datetime.now().strftime("%H:%M:%S %d/%m/%Y"),
            "file_source": os.path.basename(file_path),
            "mtime": current_mtime,
            "total_rows": len(records),
            "company": "Công ty Cổ phần Beta Solutions",
            "brand_guideline": {
                "name": "Công ty Cổ phần Beta Solutions",
                "budget_color": "#7C3AED",       # Tím Hoàng Gia
                "spent_color": "#06B6D4",        # Xanh Ngọc
                "over_budget_color": "#F43F5E",  # Hồng San Hô
                "saving_color": "#14B8A6",       # Xanh Bạc Hà
                "dark_background": "#1E1E2E",    # Xám Đậm
                "fonts": ["Outfit", "Roboto", "Segoe UI"]
            },
            "summary_preview": {
                "total_budget": total_ns,
                "total_spent": total_ct,
                "over_budget_count": vuot_count
            },
            "records": records
        }
        
        DATA_CACHE["file_path"] = file_path
        DATA_CACHE["last_mtime"] = current_mtime
        DATA_CACHE["records"] = records
        DATA_CACHE["cached_payload"] = payload
        print(f"[CACHE READY] Đã nạp thành công {len(records)} giao dịch thô. Tổng NS: {total_ns:,} ₫ | Vượt NS: {vuot_count} GD")

    return DATA_CACHE["cached_payload"]

class BetaDashboardHandler(BaseHTTPRequestHandler):
    """HTTP Request Handler phục vụ Web Dashboard và REST API cho client Polling."""

    def log_message(self, format, *args):
        # Ẩn bớt log truy vấn polling định kỳ 2s để console luôn thoáng đãng
        if "/api/raw_data" in self.path or "/api/data" in self.path:
            return
        super().log_message(format, *args)

    def do_GET(self):
        parsed_url = urlparse(self.path)
        path = parsed_url.path

        # 1. API Cung Cấp Dữ Liệu Thô (Raw Data Polling Endpoint)
        if path in ["/api/raw_data", "/api/data", "/api/raw"]:
            data = get_latest_raw_data()
            json_bytes = json.dumps(data, ensure_ascii=False).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.send_header("Cache-Control", "no-cache, no-store, must-revalidate")
            self.send_header("Content-Length", str(len(json_bytes)))
            self.end_headers()
            self.wfile.write(json_bytes)
            return

        # 2. Health check
        if path == "/api/health":
            resp = json.dumps({
                "status": "healthy",
                "service": "beta_solutions_bi_dashboard",
                "company": "Beta Solutions JSC",
                "time": datetime.now().isoformat()
            }).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(resp)
            return

        # 3. Phục vụ Web Dashboard HTML
        if path in ["/", "/index.html", "/dashboard"]:
            if os.path.exists(REPORT_HTML_PATH):
                with open(REPORT_HTML_PATH, "rb") as f:
                    content = f.read()
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(content)))
                self.end_headers()
                self.wfile.write(content)
                return
            else:
                self.send_response(404)
                self.send_header("Content-Type", "text/plain; charset=utf-8")
                self.end_headers()
                self.wfile.write(f"Chưa tìm thấy file dashboard tại: {REPORT_HTML_PATH}".encode("utf-8"))
                return

        # 4. Phục vụ tài nguyên tĩnh nếu có yêu cầu
        file_target = os.path.join(BASE_DIR, path.lstrip("/"))
        if os.path.exists(file_target) and os.path.isfile(file_target):
            content_type = "text/plain"
            if file_target.endswith(".css"):
                content_type = "text/css; charset=utf-8"
            elif file_target.endswith(".js"):
                content_type = "application/javascript; charset=utf-8"
            elif file_target.endswith(".json"):
                content_type = "application/json; charset=utf-8"
            elif file_target.endswith(".png"):
                content_type = "image/png"
            elif file_target.endswith(".svg"):
                content_type = "image/svg+xml"

            with open(file_target, "rb") as f:
                content = f.read()
            self.send_response(200)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
            return

        # 404 Không tìm thấy
        self.send_response(404)
        self.send_header("Content-Type", "text/plain; charset=utf-8")
        self.end_headers()
        self.wfile.write(b"404 Not Found")

def auto_open_chrome(url, delay_sec=1.2):
    """Tự động mở Google Chrome hoặc trình duyệt mặc định sau delay_sec giây."""
    def _launcher():
        time.sleep(delay_sec)
        chrome_paths = [
            r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
            os.path.expandvars(r"%LOCALAPPDATA%\Google\Chrome\Application\chrome.exe")
        ]
        opened = False
        for cp in chrome_paths:
            if os.path.exists(cp):
                try:
                    subprocess.Popen([cp, url])
                    print(f"[BROWSER] Đã mở Google Chrome tự động tại: {url}")
                    opened = True
                    break
                except Exception as ex:
                    print(f"[BROWSER WARNING] Lỗi mở Chrome trực tiếp: {ex}")
        if not opened:
            webbrowser.open(url)
            print(f"[BROWSER] Đã mở trình duyệt mặc định tại: {url}")

    thread = threading.Thread(target=_launcher, daemon=True)
    thread.start()

def is_port_in_use(port):
    """Kiểm tra xem port đã bị chiếm dụng chưa."""
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        return s.connect_ex(('127.0.0.1', port)) == 0

def find_available_port(start_port=8088, max_attempts=15):
    """Tìm port trống khả dụng bắt đầu từ start_port."""
    for p in range(start_port, start_port + max_attempts):
        if not is_port_in_use(p):
            return p
    return start_port

def run_server(port=8088, auto_open=True):
    """Khởi chạy server HTTP đa luồng tại port chỉ định."""
    print("=" * 70)
    print("   CÔNG TY CỔ PHẦN BETA SOLUTIONS - REALTIME BI BUDGET DASHBOARD")
    print("   Tiêu chuẩn: ai4a:build-dashboard-BI • Glassmorphism • Dark Mode")
    print("   API Dữ Liệu Thô • Client-Side Filtering • Polling 2s • 60fps")
    print("=" * 70)

    data_file = find_budget_data_file()
    if data_file:
        print(f"[DATA SOURCE] Nguồn Excel: {data_file}")
        init_data = get_latest_raw_data()
        sp = init_data.get("summary_preview", {})
        print(f"[AUDIT] Tổng NS: {sp.get('total_budget'):,} ₫ | Tổng Chi: {sp.get('total_spent'):,} ₫ | Vượt NS: {sp.get('over_budget_count')} giao dịch")
    else:
        print("[DATA WARNING] Chưa tìm thấy TH_ngan_sach_phong_ban.xlsx!")

    actual_port = port
    if is_port_in_use(actual_port):
        actual_port = find_available_port(port + 1)
        print(f"[PORT NOTICE] Port {port} đang bận, tự động chuyển sang port: {actual_port}")

    server_address = ('0.0.0.0', actual_port)
    try:
        httpd = ThreadingHTTPServer(server_address, BetaDashboardHandler)
    except OSError as e:
        print(f"[SERVER ERROR] Không thể bind port {actual_port}: {e}")
        return

    dashboard_url = f"http://localhost:{actual_port}"
    print(f"[SERVER ONLINE] Web Dashboard: {dashboard_url}")
    print(f"[POLLING API]   Raw Data API:  {dashboard_url}/api/raw_data")
    print(f"[REACTIVE]      Frontend tự xử lý lọc tức thì theo Quý & Phòng ban")
    print(f"[BRAND COLORS]  Ngân sách (#7C3AED Tím Hoàng Gia) • Chi tiêu (#06B6D4 Xanh Ngọc)")
    print(f"                Vượt NS (#F43F5E Hồng San Hô) • Tiết kiệm (#14B8A6 Xanh Bạc Hà)")
    print("=" * 70)
    print("Nhấn Ctrl+C để dừng server.")

    if auto_open:
        auto_open_chrome(dashboard_url)

    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\n[SERVER SHUTDOWN] Đang dừng server...")
        httpd.server_close()
        print("[SERVER STOPPED] Đã đóng server an toàn.")

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Beta Solutions Realtime BI Dashboard Server")
    parser.add_argument("--port", type=int, default=8088, help="Port lắng nghe (mặc định 8088)")
    parser.add_argument("--no-browser", action="store_true", help="Không tự động mở trình duyệt")
    parser.add_argument("--test-data", action="store_true", help="Chỉ kiểm tra đọc dữ liệu thô rồi thoát")

    args = parser.parse_args()

    if args.test_data:
        data = get_latest_raw_data()
        print(f"Status: {data.get('status')}")
        print(f"Total records: {data.get('total_rows')}")
        print(f"Summary Preview: {json.dumps(data.get('summary_preview'), indent=2, ensure_ascii=False)}")
        print("Mẫu 1 bản ghi đầu tiên:")
        if data.get("records"):
            print(json.dumps(data["records"][0], indent=2, ensure_ascii=False))
    else:
        run_server(port=args.port, auto_open=not args.no_browser)
